#!/usr/bin/env python3
"""生成 Board使用说明书.docx（含截图）。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "docs" / "tensorboard-overview.png"
OUT = ROOT / "Board使用说明书.docx"


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            table.rows[r].cells[c].text = str(val)
    doc.add_paragraph()


def main():
    doc = Document()
    title = doc.add_heading("Board 使用说明书", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run("2048 RL 训练实验平台 — TensorBoard 指标与操作指南\n").bold = True
    p.add_run("版本：v1.0 ｜ 日期：2026-06-07")

    doc.add_heading("1. 快速启动", level=1)
    doc.add_paragraph("在项目根目录执行：")
    for cmd in ["make board", "# 浏览器打开 http://localhost:6006", "make board PORT=6007  # 端口占用时"]:
        doc.add_paragraph(cmd, style="List Bullet")

    doc.add_heading("2. 界面总览", level=1)
    if IMG.exists():
        doc.add_picture(str(IMG), width=Inches(6.0))
        cap = doc.add_paragraph("图 1：TensorBoard 界面（Scalars 视图，多 Run 对比）")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("3. 指标说明书（9 项）", level=1)
    doc.add_paragraph("分为 train（每局结束记录）与 eval（周期性贪心评估）两组。")

    add_table(
        doc,
        ["指标 Tag", "分组", "含义", "期望趋势"],
        [
            ("train/episode_reward", "train", "单局累计 reward", "上升"),
            ("train/episode_len", "train", "单局步数", "变长=存活更久"),
            ("train/max_tile", "train", "单局最大方块", "32→512 攀升"),
            ("train/score", "train", "单局总分", "随 max_tile 升"),
            ("train/epsilon", "train", "DQN 探索率 ε", "单调下降"),
            ("train/loss", "train", "TD Loss", "逐步稳定"),
            ("eval/max_tile_median", "eval", "评估 max_tile 中位数", "≥512 达标(S2)"),
            ("eval/max_tile_mean", "eval", "评估 max_tile 均值", "辅助参考"),
            ("eval/score_mean", "eval", "评估平均分", "辅助参考"),
        ],
    )

    doc.add_heading("4. 核心指标详解", level=1)
    doc.add_heading("eval/max_tile_median（最重要）", level=2)
    doc.add_paragraph(
        "每 eval_every 步触发贪心评估，取多局 max_tile 的中位数。"
        "本项目 S2 达标线：中位数 ≥ 512。创新高时自动保存 best.pt。"
    )
    doc.add_heading("train/max_tile", level=2)
    doc.add_paragraph(
        "每局训练结束时的最大方块值，最直观反映 Agent 是否在「变聪明」。"
    )

    doc.add_heading("5. 典型操作", level=1)
    steps = [
        "左侧勾选要对比的 Run（如 quick 与 full）",
        "打开 eval/max_tile_median 卡片观察智能性趋势",
        "右侧 Smoothing 设为 0.6，开启 Ignore outliers",
        "横轴选 Step，点击曲线查看精确数值",
        "结合 runs/<id>/train.log 中的 [best] 行定位 best.pt",
    ]
    for s in steps:
        doc.add_paragraph(s, style="List Number")

    doc.add_heading("6. 目录结构", level=1)
    doc.add_paragraph(
        "runs/<时间戳>_<配置名>/\n"
        "  ├── events.out.tfevents.*  ← TensorBoard 数据源\n"
        "  ├── config.snapshot.yaml\n"
        "  ├── train.log\n"
        "  └── best.pt"
    )

    doc.add_heading("7. 常见问题", level=1)
    add_table(
        doc,
        ["问题", "解决"],
        [
            ("页面空白", "勾选左侧 Run；确认 runs/ 有 tfevents 文件"),
            ("端口占用", "make board PORT=6007"),
            ("与 GUI 区别", "TensorBoard=训练曲线；make play=对局界面"),
        ],
    )

    doc.add_paragraph("\n完整版请参阅同目录 Board使用说明书.md")

    doc.save(OUT)
    print(f"已生成: {OUT}")


if __name__ == "__main__":
    main()
